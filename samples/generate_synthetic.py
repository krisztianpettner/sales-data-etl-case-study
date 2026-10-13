#!/usr/bin/env python3
"""Generate a fully synthetic dataset in the exact layouts the ETL expects.

Every company, person, number and identifier below is invented. The generator does not
know any real value: it reads only the *layout* registry (`src/etl/config.py::FILE_CONFIG`)
and writes one workbook per configured product group × year, in that year's layout, plus
one order register per year. Then the real pipeline can run on it:

    pip install openpyxl
    python samples/generate_synthetic.py          # writes src/data/*.xlsx
    cd src && python etl/run_etl.py               # needs etl/chart.umd.min.js (Chart.js 4.4.x)

The data is deliberately messy, with the same KINDS of error the real sources had
(see docs/en/01-the-mess.md): name variants, accents, legal-form suffixes, notes typed into
the name field, decorated and accent-free statuses, '#REF!' / '-' markers, line-item rows,
won deals without a work number, missing countries, stale open quotes.
Serial numbers come from the 900+ block, so they can never coincide with a real one.
"""
import random
import sys
from datetime import date, timedelta
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent / "src"
sys.path.insert(0, str(ROOT / "etl"))
import config  # noqa: E402  — the layout registry of the exported code

rnd = random.Random(7)
AS_OF = date(2026, 10, 13)

# ── synthetic partners: (canonical name, country) ───────────────────────────────────────────
COUNTRIES = ["Magyarország"] * 6 + ["Németország", "Ausztria", "Lengyelország", "Csehország",
                                    "Románia", "Hollandia", "Svájc", "Franciaország"]
PARTNERS = [(f"Cég {i:03d}", rnd.choice(COUNTRIES)) for i in range(1, 251)]
SUFFIX = {"Magyarország": "Kft.", "Németország": "GmbH", "Ausztria": "GmbH", "Lengyelország": "Sp. z o.o.",
          "Csehország": "s.r.o.", "Románia": "S.R.L.", "Hollandia": "B.V.", "Svájc": "AG", "Franciaország": "SAS"}


def spelled(name, country):
    """One way a person might type this company's name (the A-class errors of chapter 01)."""
    suf = SUFFIX[country]
    return rnd.choices([
        f"{name} {suf}",                                  # canonical
        f"{name} {suf.lower().rstrip('.')}",              # A1: case + missing dot
        f"{name.upper()}",                                # A1: no suffix, upper case
        f"{name.replace('Cég', 'Ceg')} {suf}",            # A2: accent dropped
        f"{name.replace(' ', '-')}",                      # A2: hyphen
        f"{name} {suf} (B telephely)",                    # A4: note typed into the name field
        f"{name}  {suf} ",                                # stray whitespace
    ], weights=[50, 10, 8, 10, 6, 4, 12])[0]


# ── statuses (raw spellings, as typed) ──────────────────────────────────────────────────────
ST_TG1 = {"open": ["kiajánlott (élő)", "Kiajánlott (élő) "], "prep": ["árajánlatra vár", ">kiküldésre vár<"],
          "won": ["megnyertük", "Megnyertük", "megnyertuk"], "lost_c": ["nem mi nyertük (elveszített)"],
          "lost_u": ["vevő, nem nyerte el"], "post": ["vevő, elhalasztotta"], "noq": ["nem készült árajánlat"],
          "cnr": ["visszajelzés nélküli, így utólag lezártuk"]}
ST_TG2 = {"open": ["kiajánlott"], "prep": ["árajánlatra vár", ">>>kiküldésre vár<<<"], "won": ["megnyertük"],
          "lost_c": ["nem mi nyertük"], "lost_u": ["vevő, nem nyerte el"], "post": ["vevő, elhalasztotta"],
          "noq": ["nem készült árajánlat"], "cnr": ["visszajelzés nélküli, így utólag lezártuk"]}
LINE_MARK = "[rendelés tételsora]"


def outcome(req, pl):
    """Older requests are mostly decided; recent ones mostly open. Some old ones stay open (stale)."""
    age = (AS_OF - req).days
    w_open = 70 if age < 120 else (25 if age < 365 else 6)
    keys = ["open", "prep", "won", "lost_c", "lost_u", "post", "noq", "cnr"]
    weights = [w_open, 4, 18, 12, 8, 4, 10 if pl == "TG1" else 3, 3 if pl == "TG1" else 25]
    return rnd.choices(keys, weights=weights)[0]


def rand_date(year):
    end = AS_OF if year == AS_OF.year else date(year, 12, 31)
    return date(year, 1, 1) + timedelta(days=rnd.randrange((end - date(year, 1, 1)).days + 1))


def write_row(ws, r, cols, rec):
    for field, col in cols.items():
        if field in rec:
            ws.cell(row=r, column=col, value=rec[field])


# ── quotes ──────────────────────────────────────────────────────────────────────────────────
won_register = []          # (pl, year, work_number, customer, value) → order register


def gen_quotes(key, cfg):
    pl, year = key.split("_")[0], int(key.split("_")[1])
    yy = year % 100
    cols, start = cfg["cols"], cfg["start_row"]
    sheets = cfg["sheet"] if isinstance(cfg["sheet"], list) else [cfg["sheet"]]
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    wss = [wb.create_sheet(s) for s in sheets]
    n = 70 if pl == "TG1" else 140
    if year == AS_OF.year:
        n = int(n * 0.8)
    row = {s: start for s in sheets}
    serial = 900
    for i in range(1, n + 1):
        name, country = rnd.choice(PARTNERS)
        req = rand_date(year)
        ws = wss[0] if len(wss) == 1 or req.month <= 6 else wss[1]
        st = outcome(req, pl)
        raw = rnd.choice((ST_TG1 if pl == "TG1" else ST_TG2)[st])
        sent = req + timedelta(days=rnd.randint(1, 25)) if st not in ("prep", "noq") else None
        if rnd.random() < 0.02 and sent:
            sent = req + timedelta(days=rnd.randint(95, 160))            # E6: slow reaction
        qid = f"Q{pl[-1]}-{yy:02d}-{i:03d}"
        rec = {"quote_id": qid, "partner_name": spelled(name, country), "request_date": req,
               "sent_date": sent, "status": raw, "registry_number": f"K9{i:03d}/{yy:02d}" if sent else "-",
               "handler_name": f"Kapcsolattartó {rnd.randint(1, 40):02d}", "subject": "-",
               "revision": "REV0" if pl == "TG1" else 0, "row_number": i}
        if rnd.random() < 0.01:
            rec["request_date"] = None                                    # E6: missing date
        if "country_name" in cols:
            rec["country_name"] = country if rnd.random() > 0.15 else "#REF!"   # F1: missing / broken
            rec["country_domestic"] = "belföldi" if country == "Magyarország" else "külföldi"
        if "country" in cols:
            rec["country"] = country if rnd.random() > 0.1 else None
        work = None
        if st == "won" and rnd.random() > 0.08:                           # E5: some won without work no.
            serial += 1
            work = f"{yy:02d}{serial:03d}" if pl == "TG1" else f"W{serial:03d}{yy:02d}"
            rec["work_number"] = work
        if pl == "TG1":
            value = round(rnd.lognormvariate(11.2, 1.1), -2)
            rec["total_value_eur"] = value if st != "noq" or rnd.random() < 0.3 else None
            write_row(ws, row[ws.title], cols, rec)
            row[ws.title] += 1
        else:
            lines = rnd.choices([1, 2, 3, 4], weights=[55, 25, 12, 8])[0]
            value = 0
            for pos in range(1, lines + 1):
                qty = rnd.randint(1, 8)
                unit = round(rnd.lognormvariate(6.6, 0.8), 1)
                lrec = dict(rec) if pos == 1 else {"quote_id": qid, "status": LINE_MARK}
                lrec.update({"position": pos, "quantity": qty, "unit_price_eur": unit,
                             "line_value_eur": round(qty * unit, 1)})
                value += qty * unit
                write_row(ws, row[ws.title], cols, lrec)
                row[ws.title] += 1
        if work:
            won_register.append((pl, year, work, rec["partner_name"], value))
    # partner sheet in that year's partner layout, if the registry expects one
    if "partner_sheet" in cfg:
        ps = wb.create_sheet(cfg["partner_sheet"])
        pr = cfg.get("partner_start", 4)
        for name, country in PARTNERS:
            write_row(ps, pr, cfg["partner_cols"], {"name": f"{name} {SUFFIX[country]}", "country": country})
            pr += 1
    if "contact_sheet" in cfg:
        wb.create_sheet(cfg["contact_sheet"])
    out = config.DATA_DIR / Path(cfg["file"]).name
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


# ── order register: one workbook per year, 'won' quotes + a third division + own production ─
def gen_register():
    sys.path.insert(0, str(ROOT / "scripts"))
    import extract_work_orders as wo  # noqa: E402  — sheet name, columns, file years
    by_year = {}
    for pl, year, work, customer, value in won_register:
        by_year.setdefault(year, []).append((pl, work, customer, value))
    outs = []
    for fy in wo.FILE_YEARS:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = wo.SHEET
        r = 3
        section = None
        rows = sorted(by_year.get(fy, []), key=lambda x: x[0])
        for pl, work, customer, value in rows:
            marker = "E" if pl == "TG1" else "TG2"
            vals = [None] * 17
            vals[wo.C_MSZ] = work
            vals[wo.C_TYPE] = marker if marker != section else None      # forward-filled type column
            section = marker
            vals[wo.C_CUST] = customer
            vals[wo.C_K_EUR] = round(value * rnd.uniform(0.9, 1.05), 0)
            for c, v in enumerate(vals, 1):
                ws.cell(row=r, column=c, value=v)
            r += 1
        for k in range(1, 6):                                            # third division (4-digit numbers)
            ws.cell(row=r, column=wo.C_MSZ + 1, value=f"{fy % 100:02d}9{k}")
            ws.cell(row=r, column=wo.C_TYPE + 1, value="S" if k == 1 else None)
            ws.cell(row=r, column=wo.C_CUST + 1, value=f"Cég {rnd.randint(1, 250):03d} Kft.")
            ws.cell(row=r, column=wo.C_K_EUR + 1, value=round(rnd.lognormvariate(10.5, 0.6), 0))
            r += 1
        out = config.DATA_DIR / f"rendelesi_regiszter_{fy}.xlsx"
        wb.save(out)
        outs.append(out)
    return outs


def main():
    written = [gen_quotes(k, cfg) for k, cfg in config.FILE_CONFIG.items()]
    written += gen_register()
    for p in written:
        print("wrote", p.relative_to(ROOT.parent))


if __name__ == "__main__":
    main()
